import pdfplumber
from pptx import Presentation


class DocumentParser:

    @staticmethod
    def extract_pdf(file):
        """
        Extract text from a PDF file.
        """

        extracted_text = []

        with pdfplumber.open(file) as pdf:

            for page in pdf.pages:

                text = page.extract_text()

                if text:
                    extracted_text.append(text)

        return "\n".join(extracted_text)


    @staticmethod
    def extract_pptx(file):
        """
        Extract text from a PPTX file.
        """

        extracted_text = []

        presentation = Presentation(file)

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        extracted_text.append(text)

        return "\n".join(extracted_text)


    @staticmethod
    def extract_text(file):

        filename = file.filename.lower()

        if filename.endswith(".pdf"):

            return DocumentParser.extract_pdf(file)

        elif filename.endswith(".pptx"):

            return DocumentParser.extract_pptx(file)

        else:

            raise ValueError(
                "Unsupported file format. Please upload a PDF or PPTX file."
            )